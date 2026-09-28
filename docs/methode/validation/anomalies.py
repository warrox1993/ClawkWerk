"""Recense les écarts des formules des outils officiels par rapport à leur
méthode documentée (moyenne par sous-catégorie, puis par catégorie, N/A = seuil)."""
import openpyxl, os, re, warnings, sys
warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from officiel import FICHIERS, FONCTIONS
cellre = re.compile(r"\$?([A-Z]{1,2})\$?(\d+)(?::\$?([A-Z]{1,2})\$?(\d+))?")

def refs_col(fml, colnum):
    """Lignes référencées dans une colonne donnée, plages A1:A9 développées."""
    out = set()
    for c1, r1, c2, r2 in cellre.findall(fml):
        if col(c1) != colnum:
            continue
        out.update(range(int(r1), int(r2) + 1) if r2 else [int(r1)])
    return out
def col(c):
    n = 0
    for ch in c: n = n * 26 + ord(ch) - 64
    return n
for niv, f in FICHIERS.items():
    wb = openpyxl.load_workbook(f)
    basic = niv == "basic"
    cD, cI = (6, 7) if basic else (7, 8)     # notes
    cSD, cSI, cCD, cCI = (8, 9, 10, 11) if basic else (9, 10, 11, 12)
    creq = 5 if basic else 6
    anomalies = []
    for fn in FONCTIONS:
        ws = wb[fn]
        rows = {}; sc = ""; cat = ""
        for r in range(3, ws.max_row + 1):
            if ws.cell(r, 1).value: cat = re.search(r"\(([A-Z]{2}\.[A-Z]{2})\)", str(ws.cell(r, 1).value)).group(1)
            if ws.cell(r, 4).value: sc = str(ws.cell(r, 4).value).split(":")[0].strip()
            v = ws.cell(r, creq).value
            if v and ":" in str(v): rows[r] = (str(v).split(":")[0].strip(), sc, cat)
        # groupes de sous-catégorie : formules de la colonne doc/impl de sous-catégorie
        groups = {}  # ligne de la formule -> lignes de notes référencées (doc) , substitution N/A ?
        for r in range(3, ws.max_row + 1):
            for cc, axis, src in ((cSD, "doc", cD), (cSI, "impl", cI)):
                fml = ws.cell(r, cc).value
                if not isinstance(fml, str) or not fml.startswith("="): continue
                refrows = sorted(refs_col(fml, src))
                na = fml.count('"N/A"')
                groups[(r, axis)] = refrows
                if na == 0:
                    anomalies.append(f"{fn}!{openpyxl.utils.get_column_letter(cc)}{r} ({rows.get(r,('?',))[0]}, {axis}) : pas de substitution N/A ({fml})")
                scs = {rows[x][1] for x in refrows if x in rows}
                if len(scs) > 1:
                    anomalies.append(f"{fn}!{openpyxl.utils.get_column_letter(cc)}{r} ({axis}) : moyenne de plusieurs sous-catégories {sorted(scs)} = {[rows[x][0] for x in refrows]}")
        # couverture : chaque exigence dans exactement un groupe par axe
        for axis in ("doc", "impl"):
            cov = {}
            for (r, a), refs in groups.items():
                if a == axis:
                    for x in refs: cov.setdefault(x, []).append(r)
            for r, (rid, sc, cat) in rows.items():
                if len(cov.get(r, [])) != 1:
                    anomalies.append(f"{fn} {rid} ({axis}) : présente dans {len(cov.get(r, []))} groupe(s)")
        # une sous-catégorie éclatée en plusieurs groupes (chaque exigence
        # comptée comme une sous-catégorie)
        for axis in ("doc", "impl"):
            parsc = {}
            for (g, a), refs in groups.items():
                if a == axis and g in rows:
                    parsc.setdefault(rows[g][1], []).append(g)
            for scid, gs in parsc.items():
                if len(gs) > 1:
                    anomalies.append(f"{fn} {scid} ({axis}) : sous-catégorie éclatée en {len(gs)} groupes (lignes {sorted(gs)})")
        # catégories : K/L doivent référencer tous les groupes de la catégorie
        for r in range(3, ws.max_row + 1):
            for cc, axis, src in ((cCD, "doc", cSD), (cCI, "impl", cSI)):
                fml = ws.cell(r, cc).value
                if not isinstance(fml, str) or not fml.startswith("="): continue
                cat = rows[r][2] if r in rows else "?"
                refs = sorted(x for x in refs_col(fml, src) if (x, axis) in groups)
                expected = sorted({g for (g, a) in groups if a == axis and g in rows and rows[g][2] == cat})
                if refs != expected:
                    anomalies.append(f"{fn}!{openpyxl.utils.get_column_letter(cc)}{r} ({cat}, {axis}) : référence {refs}, groupes présents {expected}")
    print(f"===== {niv} : {len(anomalies)} anomalie(s)")
    for a in anomalies: print("  ", a)
