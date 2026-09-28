"""Lecture des outils d'auto-évaluation officiels CyFun 2025 (CCB) : liste des
exigences, Key Measures, niveau, sous-catégorie et cellules de saisie."""
import openpyxl, os, re, warnings
warnings.filterwarnings("ignore")
# Répertoire où ont été téléchargés les trois outils officiels depuis
# https://cyfun.eu/en/cyberfundamentals-framework-2025 (non redistribués ici).
D = os.environ.get("CYFUN_OUTILS", ".") + "/"
FICHIERS={"basic":D+"CyFun2025_Self-Assessment_tool_BASIC_v2026_02_20.xlsx",
          "important":D+"CyFun2025_Self-Assessment_tool_IMPORTANT_v2026_02_20.xlsx",
          "essential":D+"CyFun2025_Self-Assessment_tool_ESSENTIAL_v3.1.xlsx"}
FONCTIONS=["GOVERN","IDENTIFY","PROTECT","DETECT","RESPOND","RECOVER"]
def exigences(niveau):
    wb=openpyxl.load_workbook(FICHIERS[niveau])
    basic = niveau=="basic"
    creq = 5 if basic else 6   # colonne E ou F
    cdoc = 6 if basic else 7
    out=[]
    for f in FONCTIONS:
        ws=wb[f]
        sc_courante=''
        for r in range(3, ws.max_row+1):
            scv=ws.cell(r,4).value
            if scv: sc_courante=str(scv).split(':')[0].strip()
            v=ws.cell(r,creq).value
            if not v or not isinstance(v,str) or ':' not in v: continue
            rid,txt=v.split(':',1)
            out.append(dict(fonction=f, ligne=r, id=rid.strip(), texte=" ".join(txt.split()),
                km=(str(ws.cell(r,3).value or '').strip()=="Key Measure"),
                niveau=("Basic" if basic else str(ws.cell(r,5).value).strip()),
                sous_cat=sc_courante, cat=sc_courante.split('-')[0],
                col_doc=cdoc, col_impl=cdoc+1))
    return out
